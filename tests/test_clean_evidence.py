"""Regression checks on the compact full-corpus research evidence."""
import hashlib
import json
from pathlib import Path
import re

import numpy as np
import pandas as pd
from pptx import Presentation
import pymupdf

ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "artifacts/clean_text_full"
EVIDENCE = ARTIFACTS / "evidence"


def test_completed_full_run_not_the_old_sample():
    run = json.loads((EVIDENCE / "experiment.json").read_text())
    prep = json.loads((EVIDENCE / "preparation.json").read_text())
    assert run['status'] == 'completed'
    assert run['units'] == prep['model_units'] - run['vocabulary_exclusions']
    assert run['units'] > 200 and run['documents'] == prep['included_documents']
    assert prep['input_format'] == 'cleaned TXT only'
    encoder = json.loads((EVIDENCE / 'encoder_metadata.json').read_text())
    assert encoder['embedding_shape'] == [run['units'], 768]
    assert encoder['total_chunks'] >= run['units']


def test_every_source_is_accounted_for_without_content_changes():
    table = pd.read_csv(EVIDENCE / 'file_audit.csv', keep_default_na=False)
    prep = json.loads((EVIDENCE / 'preparation.json').read_text())
    assert len(table) == prep['source_files']
    assert table.path.nunique() == len(table)
    root = ROOT / 'data/processed/clean_text'
    for row in table.itertuples():
        text = (root / row.path).read_text(encoding='utf-8-sig')
        digest = hashlib.sha256(re.sub(r'\s+', ' ', text).strip().encode()).hexdigest()
        assert digest == row.content_sha256
    assert table.content_sha256.nunique() == prep['content_groups']
    assert len(table) - table.content_sha256.nunique() == prep['duplicate_copies']


def test_metrics_and_temporal_denominators_are_consistent():
    trials = pd.read_csv(EVIDENCE / 'tuning_results.csv')
    yearly = pd.read_csv(EVIDENCE / 'modeled_yearly_counts.csv').set_index('year')
    assert len(trials) == 4
    assert set(yearly.index) == set(range(2016, 2026))
    for model in ['lexical', 'semantic']:
        metrics = json.loads((EVIDENCE / model / 'metrics.json').read_text())
        initial = trials.query("trial == 'initial' and model == @model").iloc[0]
        for key in ['cv_coherence', 'topic_diversity', 'topic_coverage', 'stability_ari_mean']:
            assert np.isclose(metrics[key], initial[key])
        assert np.isclose(metrics['topic_coverage'] + metrics['outlier_fraction'], 1)
        shares = pd.read_csv(EVIDENCE / model / 'yearly_prevalence.csv')
        assert np.allclose(shares.groupby('year').section_share.sum(), 1)
        assert shares.groupby('year').section_count.sum().equals(yearly.units)
        weighted = pd.read_csv(EVIDENCE / model / 'document_weighted_prevalence.csv')
        assert np.allclose(weighted.groupby('year').share.sum(), 1)
        assert np.array_equal(weighted.groupby('year').documents_in_year.first().values, yearly.documents.values)


def test_deck_contains_no_legacy_archival_workflow():
    path = ARTIFACTS / 'Clean_Text_Full_Corpus_Initial_Results'
    deck = Presentation(str(path.with_suffix('.pptx')))
    assert len(deck.slides) == 27
    text = '\n'.join(shape.text for slide in deck.slides for shape in slide.shapes if shape.has_text_frame)
    assert 'obsidian' not in text.lower()
    with pymupdf.open(path.with_suffix('.pdf')) as pdf:
        assert len(pdf) == len(deck.slides)
        assert all(page.get_text().strip() for page in pdf)
    assert len(list((ARTIFACTS / 'figures').glob('year_*_comparison.png'))) == 10


def test_no_fabricated_expert_scores():
    for model in ['lexical', 'semantic']:
        table = pd.read_csv(EVIDENCE / model / 'expert_review_template.csv', keep_default_na=False)
        scores = [column for column in table if column.endswith('_1_to_5')]
        assert len(scores) == 5
        assert table[scores].eq('').all().all()
