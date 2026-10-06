"""Check the committed real-pilot report against its recorded measurements."""
import hashlib
import json
from pathlib import Path
import sys

import numpy as np
import pandas as pd
from pptx import Presentation

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = ROOT / "artifacts/pilot/evidence"
sys.path.insert(0, str(ROOT / "scripts"))
from run_comparison import prepare_sample


def test_completed_experiment_and_reproducible_sample():
    manifest = json.loads((EVIDENCE / "experiment.json").read_text())
    assert manifest["status"] == "completed"
    config = manifest["config"]
    source = ROOT / config["source"]
    portable = json.loads((EVIDENCE / "portable_source.json").read_text())
    assert hashlib.sha256(source.read_bytes().replace(b"\r\n", b"\n")).hexdigest() == portable["source_sha256_lf"]
    sample, _ = prepare_sample(source, config)
    assert len(sample) == manifest["sections"] == 200
    fingerprint = hashlib.sha256(json.dumps(sample.text.tolist(), ensure_ascii=False).encode()).hexdigest()
    metadata = json.loads((EVIDENCE / "semantic/model_metadata.json").read_text())
    assert fingerprint == metadata["ordered_text_sha256"]
    assert sample.verification_status.ne("verified").sum() == manifest["unverified_sections"]


def test_reported_metrics_match_initial_trial():
    table = pd.read_csv(EVIDENCE / "tuning_results.csv")
    assert set(table.trial) == {"initial", "conservative"}
    assert len(table) == 4
    for name in ["lexical", "semantic"]:
        metrics = json.loads((EVIDENCE / name / "metrics.json").read_text())
        row = table.query("trial == 'initial' and model == @name").iloc[0]
        for key in ["cv_coherence", "topic_diversity", "topic_coverage", "stability_ari_mean"]:
            assert np.isclose(row[key], metrics[key])
        assert np.isclose(metrics["topic_coverage"] + metrics["outlier_fraction"], 1)
        assert np.isclose(metrics["outlier_fraction"], metrics["outlier_sections"] / metrics["sections"])


def test_yearly_denominators_include_every_section():
    for name in ["lexical", "semantic"]:
        table = pd.read_csv(EVIDENCE / name / "yearly_prevalence.csv")
        assert set(table.year) == set(range(2016, 2026))
        assert table.groupby("year").section_count.sum().eq(20).all()
        assert np.allclose(table.groupby("year").section_share.sum(), 1)
        assert (table.total_sections_in_year == 20).all()


def test_presentation_and_pdf_are_complete():
    import pymupdf
    path = ROOT / "artifacts/pilot/Ordinance_Two_Model_Initial_Results"
    presentation = Presentation(str(path.with_suffix(".pptx")))
    assert len(presentation.slides) == 28
    with pymupdf.open(path.with_suffix(".pdf")) as pdf:
        assert len(pdf) == len(presentation.slides)
        assert all(page.get_text().strip() for page in pdf)
    assert len(list((ROOT / "artifacts/pilot/figures").glob("year_*_comparison.png"))) == 10


def test_expert_reviews_are_blank_not_fabricated():
    for name in ["lexical", "semantic"]:
        table = pd.read_csv(EVIDENCE / name / "expert_review_template.csv", keep_default_na=False)
        scores = [c for c in table if c.endswith("_1_to_5")]
        assert len(scores) == 5
        assert table[scores].eq("").all().all()
