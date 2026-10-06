"""Build an editable PPTX and PNG evidence only from completed real experiments."""
from __future__ import annotations

import argparse
import hashlib
import importlib.metadata
import json
from pathlib import Path
import shutil
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
NAVY, TEAL, ORANGE, INK, GRAY = "10243A", "087F8C", "D47732", "18334A", "607487"
COLORS = {"lexical": "#" + TEAL, "semantic": "#" + ORANGE}
NAMES = {"lexical": "Model A · Lexical", "semantic": "Model B · Legal-BERT"}
plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11, "axes.spines.top": False,
                     "axes.spines.right": False, "axes.labelcolor": "#" + INK,
                     "text.color": "#" + INK, "figure.facecolor": "white"})


def pretty(value, digits=3):
    return "Not measured" if pd.isna(value) else f"{value:.{digits}f}"


def save_table(df, path, title, note=""):
    fig, ax = plt.subplots(figsize=(14, max(3.5, len(df) * .52 + 1.7)))
    ax.axis("off")
    ax.set_title(title, loc="left", fontsize=19, fontweight="bold", pad=20)
    table = ax.table(cellText=df.astype(str).values, colLabels=df.columns, loc="center", cellLoc="center")
    table.auto_set_font_size(False)
    table.set_fontsize(10)
    table.scale(1, 2)
    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor("white")
        cell.set_facecolor("#" + NAVY if row == 0 else ("#EEF4F7" if row % 2 else "#F8FAFC"))
        if row == 0:
            cell.set_text_props(color="white", weight="bold")
    fig.text(.02, .015, note, fontsize=9, color="#" + GRAY)
    fig.savefig(path, dpi=170, bbox_inches="tight")
    plt.close(fig)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
    return box


def picture(slide, path, x=0.7, y=1.65, w=11.9, h=4.85):
    from PIL import Image
    with Image.open(path) as image:
        ratio = image.width / image.height
    width, height = min(w, h * ratio), min(h, w / ratio)
    slide.shapes.add_picture(str(path), Inches(x + (w - width) / 2), Inches(y + (h - height) / 2),
                             width=Inches(width), height=Inches(height))


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--experiment", type=Path, default=ROOT / "results/experiments/pilot")
    parser.add_argument("--output", type=Path, default=ROOT / "artifacts/pilot")
    args = parser.parse_args()
    run, out = args.experiment.resolve(), args.output.resolve()
    manifest = json.loads((run / "experiment.json").read_text(encoding="utf-8"))
    if manifest["status"] != "completed":
        raise ValueError("Refusing to present an incomplete experiment as completed results")
    if out.exists():
        raise ValueError("Presentation output exists; choose a new --output to preserve evidence")
    figures, evidence = out / "figures", out / "evidence"
    figures.mkdir(parents=True)
    evidence.mkdir()
    config = manifest["config"]
    if hashlib.sha256((ROOT / config["source"]).read_bytes()).hexdigest() != manifest["source_sha256"]:
        raise ValueError("Source corpus changed since this experiment; do not mix evidence")
    source_lf = (ROOT / config["source"]).read_bytes().replace(b"\r\n", b"\n")
    (evidence / "portable_source.json").write_text(json.dumps({
        "source_sha256_lf": hashlib.sha256(source_lf).hexdigest(),
        "note": "CRLF normalized to LF only, for cross-platform Git checkout verification. Original byte hash remains in experiment.json."
    }, indent=2), encoding="utf-8")
    (evidence / "render_manifest.json").write_text(json.dumps({
        "generator_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
        "packages": {name: importlib.metadata.version(name) for name in ["python-pptx", "matplotlib", "numpy", "pandas"]},
        "experiment_sha256": hashlib.sha256((run / "experiment.json").read_bytes()).hexdigest()
    }, indent=2), encoding="utf-8")
    trial = config["trials"][0]["name"]
    initial = run / trial
    metrics = pd.read_csv(run / "tuning_results.csv")
    audit = pd.read_csv(run / "corpus_audit.csv")
    models = ["lexical", "semantic"]
    assignments, keywords = {}, {}
    for name in models:
        assignments[name] = pd.read_csv(initial / name / "section_topics.csv", keep_default_na=False)
        keywords[name] = pd.read_csv(initial / name / "topic_keywords.csv")
        dest = evidence / name
        dest.mkdir()
        for file in ["metrics.json", "model_metadata.json", "yearly_prevalence.csv", "topic_keywords.csv",
                     "expert_review_template.csv"]:
            shutil.copy2(initial / name / file, dest / file)
        review = pd.read_csv(dest / "expert_review_template.csv", keep_default_na=False)
        for column in ["reviewer_id", "label_appropriateness_1_to_5", "temporal_interpretation_1_to_5",
                       "overall_research_quality_1_to_5"]:
            review[column] = ""
        review["rating_instructions"] = "1-5 or N/A; expert completion pending; flag dimension means below 3.00"
        review.to_csv(dest / "expert_review_template.csv", index=False)
    for file in ["experiment.json", "corpus_audit.csv", "tuning_results.csv"]:
        shutil.copy2(run / file, evidence / file)
    shutil.copy2(initial / "run_config.json", evidence / "initial_run_config.json")
    shutil.copy2(initial / "cross_model_agreement.json", evidence / "cross_model_agreement.json")
    # Record an auditable source for every number and image, without redistributing all section text.
    (evidence / "README.md").write_text(
        "# Measured pilot evidence\n\nGenerated from the completed experiment identified in experiment.json. "
        "Both models use the same ordered sample and shared unigram vocabulary. "
        "Trial labels are parameter sensitivity checks, not held-out model selection. "
        "Metrics are intrinsic, in-sample, section-level scores. No expert scores have been supplied.\n\n"
        "Figures are rendered directly from CSV/JSON/NPY artifacts, not fabricated UI screenshots. "
        "Yearly denominators include outliers. Topic IDs are model-specific. "
        "Raw embeddings and full text remain in the ignored local results/experiments folder.\n", encoding="utf-8")

    # Two pipeline paths; c-TF-IDF follows initial assignment because classes must exist first.
    fig, axes = plt.subplots(2, 1, figsize=(14, 5))
    paths = {
        "lexical": ["Shared sections", "TF-IDF", "SVD + L2 norm", "UMAP", "HDBSCAN", "c-TF-IDF\n+ yearly shares"],
        "semantic": ["Shared sections", "Legal-BERT\n512-token chunks", "Token-weighted\nmean + L2 norm", "UMAP", "HDBSCAN", "c-TF-IDF\n+ yearly shares"]}
    for ax, name in zip(axes, models):
        ax.axis("off")
        ax.text(0, .92, NAMES[name], color=COLORS[name], weight="bold", fontsize=15)
        for i, label in enumerate(paths[name]):
            x = .015 + i * .166
            ax.text(x + .065, .42, label, ha="center", va="center", fontsize=11,
                    bbox=dict(boxstyle="round,pad=.7", facecolor="#F0F5F8", edgecolor=COLORS[name]))
            if i < 5:
                ax.annotate("", xy=(x + .16, .42), xytext=(x + .139, .42), arrowprops=dict(arrowstyle="->"))
    fig.tight_layout()
    fig.savefig(figures / "pipelines.png", dpi=180, bbox_inches="tight")
    plt.close(fig)

    settings = pd.DataFrame([{**{k: t[k] for k in ["name", "neighbors", "min_dist", "min_cluster_size", "min_samples"]}}
                             for t in config["trials"]])
    settings.columns = ["Trial", "UMAP neighbors", "UMAP min_dist", "HDBSCAN min size", "HDBSCAN min samples"]
    save_table(settings, figures / "tuning_parameters.png", "Parameters actually executed",
               f"Shared: seed {config['seed']} | UMAP dimensions {config['common']['umap_components']} | "
               f"top words {config['common']['top_words']} | stability runs {config['common']['stability_runs']} | no BERT fine-tuning")
    selected = metrics[["trial", "model", "topics_excluding_outliers", "cv_coherence", "topic_diversity", "topic_coverage", "stability_ari_mean"]].copy()
    for col in selected.columns[3:]:
        selected[col] = selected[col].map(pretty)
    selected.columns = ["Trial", "Model", "Topics", "C_v", "Diversity", "Coverage", "Seed ARI"]
    save_table(selected, figures / "tuning_results.png", "Measured sensitivity results · same sample",
               "No held-out evaluation; no configuration is declared optimal. ARI includes the outlier label as one cluster.")
    scores = metrics[metrics.trial.eq(trial)].set_index("model")
    table = pd.DataFrame({"Metric": ["C_v coherence", "Top-10 diversity", "Assigned section share", "Outlier share", "Seed stability (ARI)", "Topics (not outliers)"],
                          **{NAMES[n]: [pretty(scores.loc[n, key]) for key in ["cv_coherence", "topic_diversity", "topic_coverage", "outlier_fraction", "stability_ari_mean", "topics_excluding_outliers"]] for n in models}})
    for name in models:
        table.loc[table["Metric"].eq("Topics (not outliers)"), NAMES[name]] = str(int(scores.loc[name, "topics_excluding_outliers"]))
    save_table(table, figures / "initial_metrics.png", "Initial configuration · measured model comparison",
               "Higher intrinsic scores do not establish legal validity. No accuracy, F1 or perplexity is claimed.")

    # Actual vector trace and both models' source-linked excerpts.
    vectors = np.load(initial / "semantic/section_vectors.npy", allow_pickle=False)
    metadata = json.loads((initial / "semantic/model_metadata.json").read_text())
    row = assignments["semantic"].iloc[0]
    snippet = {
        "ordinance_id": row.ordinance_id, "section_id": row.section_id, "year": int(row.year),
        "source_file": row.get("source_file", ""), "text_excerpt": row.text[:550],
        "semantic_topic": int(row.topic), "embedding_shape": list(vectors.shape),
        "first_8_values": vectors[0, :8].astype(float).tolist(), "l2_norm": float(np.linalg.norm(vectors[0])),
        "encoder_revision": metadata["resolved_model_revision"], "pooling": metadata["pooling"]}
    (evidence / "bert_snippet.json").write_text(json.dumps(snippet, indent=2), encoding="utf-8")
    fig, ax = plt.subplots(figsize=(14, 5))
    fig.set_facecolor("#" + NAVY)
    ax.axis("off")
    trace = (f"ENCODER  {metadata['encoder']}\n"
             f"REVISION {metadata['resolved_model_revision']}\n"
             f"SOURCE   {row.ordinance_id} / {row.section_id} / year {row.year}\n"
             f"OUTPUT   {vectors.shape[0]} sections x {vectors.shape[1]} dimensions\n"
             f"VECTOR   {np.array2string(vectors[0, :8], precision=5)} ...\n"
             f"L2 NORM  {np.linalg.norm(vectors[0]):.6f}   |   ASSIGNED TOPIC {row.topic}\n\n"
             + textwrap.fill(row.text[:320], 100))
    ax.text(.015, .95, trace, va="top", color="#EAF4FA", fontfamily="DejaVu Sans Mono", fontsize=12, linespacing=1.7)
    fig.savefig(figures / "bert_output_snippet.png", dpi=160, bbox_inches="tight", facecolor=fig.get_facecolor())
    plt.close(fig)

    def topic_label(name, topic, words=3):
        if topic == -1:
            return "Outliers (-1)"
        terms = keywords[name].query("topic == @topic").sort_values("rank").term.head(words).tolist()
        return f"T{topic}: " + ", ".join(terms)

    fig, ax = plt.subplots(figsize=(13, 4.5))
    for name in models:
        frame = assignments[name]
        coverage = frame.assign(assigned=frame.topic.ne(-1)).groupby("year").assigned.mean()
        ax.plot(coverage.index, coverage.values * 100, "o-", label=NAMES[name], color=COLORS[name], linewidth=2)
    ax.set(ylim=(0, 105), ylabel="Assigned sections / sampled sections (%)", xlabel="Enactment year (source metadata)")
    ax.set_xticks(config["years"])
    ax.grid(axis="y", alpha=.2)
    ax.legend(loc="lower left")
    fig.tight_layout()
    fig.savefig(figures / "yearly_coverage.png", dpi=170)
    plt.close(fig)

    for name in models:
        frame = assignments[name]
        top = frame.loc[frame.topic.ne(-1), "topic"].value_counts().head(8).index.tolist()
        if not top:
            top = [-1]
        prevalence = pd.read_csv(initial / name / "yearly_prevalence.csv")
        matrix = prevalence.pivot(index="topic", columns="year", values="section_share").reindex(index=top, columns=config["years"], fill_value=0)
        fig, ax = plt.subplots(figsize=(13, 5))
        image = ax.imshow(matrix.to_numpy() * 100, aspect="auto", cmap="Blues" if name == "lexical" else "Oranges", vmin=0, vmax=100)
        ax.set_yticks(range(len(top)), [topic_label(name, t) for t in top], fontsize=10)
        ax.set_xticks(range(len(config["years"])), config["years"])
        for i in range(len(top)):
            for j in range(len(config["years"])):
                ax.text(j, i, f"{matrix.iloc[i, j] * 100:.0f}", ha="center", va="center", fontsize=9,
                        color="white" if matrix.iloc[i, j] > .5 else "#" + INK)
        fig.colorbar(image, ax=ax, label="Share of all sampled sections in year (%)")
        fig.tight_layout()
        fig.savefig(figures / f"{name}_yearly_heatmap.png", dpi=170)
        plt.close(fig)

    for year in config["years"]:
        fig, axes = plt.subplots(1, 2, figsize=(14, 5.3))
        for ax, name in zip(axes, models):
            frame = assignments[name].query("year == @year")
            counts = frame.topic.value_counts()
            top = counts[counts.index != -1].head(3)
            labels = [topic_label(name, t) for t in top.index]
            values = top.tolist()
            other = int(counts[counts.index != -1].sum() - top.sum())
            if other:
                labels.append("Other assigned topics")
                values.append(other)
            labels.append("Outliers (-1)")
            values.append(int(counts.get(-1, 0)))
            ax.barh(range(len(labels)), np.array(values) / len(frame) * 100, color=COLORS[name], alpha=.85)
            ax.set_yticks(range(len(labels)), [textwrap.fill(x, 27) for x in labels], fontsize=10)
            ax.invert_yaxis()
            ax.set_xlim(0, 110)
            ax.set_xlabel("Share of sampled sections (%)")
            ax.set_title(f"{NAMES[name]}\n{len(frame)} sections / {frame.ordinance_id.nunique()} ordinances", fontsize=13, fontweight="bold")
            for i, value in enumerate(values):
                ax.text(value / len(frame) * 100 + 1, i, f"{value}/{len(frame)}", va="center", fontsize=10)
        fig.suptitle(f"{year} · exploratory within-year comparison", x=.02, ha="left", fontsize=19, fontweight="bold")
        fig.tight_layout(rect=[0, 0, 1, .91], w_pad=3)
        fig.savefig(figures / f"year_{year}_comparison.png", dpi=170)
        plt.close(fig)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    def slide(title, section="INITIAL RESULTS", note=None):
        s = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(s, section.upper(), .65, .24, 12, .35, 11, TEAL, True)
        add_text(s, title, .65, .75, 12.1, .75, 30, INK, True)
        bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(.7), Inches(1.48), Inches(.75), Inches(.045))
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGBColor.from_string(TEAL)
        bar.line.fill.background()
        add_text(s, "CORPUZ · LAMSIN · NURIL    /    EXPLORATORY PILOT — NOT VALIDATED THESIS FINDINGS", .65, 7.03, 11.8, .23, 9, GRAY)
        add_text(s, f"{len(prs.slides):02}", 12.15, 7, .55, .3, 10, GRAY)
        if note:
            s.notes_slide.notes_text_frame.text = note
        return s

    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = RGBColor.from_string(NAVY)
    add_text(s, "DAVAO CITY MUNICIPAL ORDINANCES", .85, .8, 11.5, .5, 15, "63D5D0", True)
    add_text(s, "Two models.\nOne shared ordinance sample.", .85, 1.8, 11.5, 2, 40, "FFFFFF", True)
    add_text(s, "Lexical TF-IDF + c-TF-IDF  vs  Legal-BERT + BERTopic", .85, 4.1, 11.5, .65, 21, "DBE8F2")
    add_text(s, f"Initial measured results · {manifest['sections']} sections · 2016–2025\nCorpuz · Lamsin · Nuril", .85, 5.2, 11.5, 1, 18, "DBE8F2")
    add_text(s, "EXPLORATORY / UNVERIFIED OCR / NOT A FINAL MODEL SELECTION", .85, 6.9, 11.5, .3, 12, "F9BF88", True)

    s = slide("What this experiment can—and cannot—show", "RESEARCH STATUS")
    add_text(s, "MEASURED NOW\nTwo complete model paths on identical sections\nTwo parameter configurations with three seed runs\nC_v, diversity, coverage, outliers and stability\nSource-linked BERT vectors and yearly visualizations", .7, 1.85, 6, 4.6, 20)
    add_text(s, "NOT YET ESTABLISHED\nHuman-verified corpus or legal topic labels\nHeld-out generalization or optimal parameters\nExpert validation and final thesis conclusions\nPopulation-level changes in municipal policy", 7.05, 1.85, 5.6, 4.6, 20, GRAY)

    s = slide("Newer methodology governs the implementation", "SOURCE ALIGNMENT")
    add_text(s, "Milestone-1_Corpus-Lamsin-Nuril.pdf\n2016–2025; section-level Legal-BERT, UMAP, HDBSCAN, c-TF-IDF; temporal analysis and expert validation.\n\ndiagram.png\nExplicit Model A lexical path versus Model B embedding path.\n\nThesis_Corpuz, Lamsin, Nuril.pdf\nEarlier LDA proposal. Retained as historical context, not the current two-model specification.", .8, 1.85, 11.7, 4.8, 20)

    s = slide("Corpus readiness is the largest limitation", "DATA AUDIT")
    source = pd.read_csv(ROOT / config["source"], keep_default_na=False)
    unverified = int(source.verification_status.ne("verified").sum())
    sampling = f"{config['sections_per_year']} seeded random sections per year" if config.get("sections_per_year") else "All eligible sections"
    add_text(s, f"{len(source):,} source sections / {source.ordinance_id.nunique():,} ordinances\n{unverified:,} source sections are unverified; {len(source) - unverified:,} verified.\n\nPilot: {manifest['sections']} sections / {manifest['ordinances']} ordinances\n{sampling}; same ordered sample for both models.\n{manifest['unverified_sections']} sampled sections remain unverified; {manifest['duplicate_text_sections']} exact repeated texts retained.\n\nNo silent OCR correction, semantic relabeling or invented expert scores.", .8, 1.8, 11.7, 4.9, 21)

    s = slide("Shared inputs; representation is the main difference", "MODEL PIPELINES")
    picture(s, figures / "pipelines.png", h=4.15)
    add_text(s, "c-TF-IDF requires existing classes: Model A first clusters TF-IDF/SVD vectors. Both paths then derive topic keywords with c-TF-IDF.", .8, 6.05, 11.7, .7, 16, GRAY)

    s = slide("The BERT module: frozen encoder, not fine-tuning", "SEMANTIC PIPELINE")
    add_text(s, "nlpaueb/legal-bert-base-uncased · pinned model revision\nLong sections → non-overlapping chunks, up to 512 tokens including special tokens.\nToken-weighted mean pooling excludes padding and special tokens.\nAggregate all chunks → L2-normalized section embedding → UMAP → HDBSCAN.\nEmbedding vectors retain section, ordinance and year associations.\n\nLegal pretraining is not proof of suitability for Philippine municipal law. No learning rate or training epochs apply: encoder weights are frozen.", .8, 1.85, 11.7, 4.9, 20)

    s = slide("Recorded parameters—not hypothetical tuning", "PARAMETER EVIDENCE")
    picture(s, figures / "tuning_parameters.png", h=3.8)
    add_text(s, "Pictures are rendered from the executed configuration. Sensitivity trials change several settings together; they do not isolate causal parameter effects.", .8, 5.95, 11.7, .8, 17, GRAY)

    s = slide("Initial model comparison", note="Source: evidence/lexical/metrics.json and evidence/semantic/metrics.json. In-sample intrinsic scores; unverified OCR pilot.")
    picture(s, figures / "initial_metrics.png")

    s = slide("What the initial numbers suggest—not prove", "COMPARATIVE READING")
    a, b = scores.loc["lexical"], scores.loc["semantic"]
    add_text(s, f"Coherence: lexical {a.cv_coherence:.3f} vs Legal-BERT {b.cv_coherence:.3f}.\nKeyword diversity: lexical {a.topic_diversity:.3f} vs Legal-BERT {b.topic_diversity:.3f}.\nAssigned sections: lexical {a.topic_coverage:.1%} vs Legal-BERT {b.topic_coverage:.1%}.\nSeed stability: lexical {a.stability_ari_mean:.3f} vs Legal-BERT {b.stability_ari_mean:.3f}.\n\nThe largest topics contain formal wording such as city, ordinance and attested. Strong intrinsic metrics may reward repeated boilerplate rather than substantive policy themes. More coverage alone does not mean better legal topics.\n\nTreat this as a working baseline for corpus cleanup—not a verdict on either model.", .8, 1.8, 11.7, 4.95, 19)

    s = slide("Do results move when clustering settings change?", "PARAMETER SENSITIVITY")
    picture(s, figures / "tuning_results.png", h=4.35)
    add_text(s, "Same input, vocabulary and frozen encoder. Cached BERT embeddings are reused after text-hash validation. Runtime for a cached trial is not an end-to-end speed comparison.", .8, 6.12, 11.7, .6, 15, GRAY)

    s = slide("A real Legal-BERT output snippet", "EXECUTION EVIDENCE")
    picture(s, figures / "bert_output_snippet.png")

    s = slide("Representative output: retain source context", "INTERPRETATION")
    excerpt_rows = []
    for i, name in enumerate(models):
        frame = assignments[name]
        chosen = frame[frame.topic.ne(-1)]
        if chosen.empty:
            add_text(s, f"{NAMES[name]}\nAll sampled sections are outliers; no substantive topic label assigned.", .8 + i * 6.2, 1.9, 5.6, 4.7, 18)
            continue
        topic = chosen.topic.value_counts().index[0]
        record = chosen[chosen.topic.eq(topic)].iloc[0]
        text = f"{NAMES[name]}\n{topic_label(name, topic, 5)}\n\n{record.ordinance_id} / {record.section_id} / {record.year}\n\n{record.text[:410]}"
        add_text(s, text, .8 + i * 6.2, 1.8, 5.55, 4.95, 16)
        excerpt_rows.append({"model": name, "topic": int(topic), "ordinance_id": record.ordinance_id,
                             "section_id": record.section_id, "year": int(record.year), "text_excerpt": record.text[:410]})
    (evidence / "representative_excerpts.json").write_text(json.dumps(excerpt_rows, indent=2), encoding="utf-8")
    add_text(s, "Keywords are machine-generated, not expert-approved policy labels. OCR artifacts and legal boilerplate remain visible.", .8, 6.45, 11.7, .45, 14, GRAY)

    s = slide("How much of each year's sample gets a topic?", "YEARLY COMPARISON")
    picture(s, figures / "yearly_coverage.png", h=4.45)
    add_text(s, "Denominator: all sampled sections in each year, including outliers. Balanced sampling does not estimate ordinance production or full-corpus topic prevalence.", .8, 6.27, 11.7, .55, 15, GRAY)
    for name in models:
        s = slide(f"{NAMES[name]} · topic shares across years", "TEMPORAL SNAPSHOT")
        picture(s, figures / f"{name}_yearly_heatmap.png", h=4.55)
        add_text(s, "Top eight global topics only; omitted topics and outliers mean columns need not sum to 100%. These are sample-level patterns, not verified policy trends.", .8, 6.35, 11.7, .55, 14, GRAY)

    s = slide("How to read the metrics", "EVALUATION CONTRACT")
    add_text(s, "C_v: sliding-window coherence of top unigram keywords (Gensim; same input texts).\nDiversity: unique topic keywords / all retained top-keyword slots; excludes outliers.\nCoverage: assigned sections / all retained sections; outlier share is its complement.\nStability: mean adjusted Rand index vs two additional UMAP seeds; outliers treated as one cluster.\n\nAll scores are intrinsic and in-sample. Accuracy/F1 need reference labels. BERTopic has no directly comparable probabilistic perplexity. Expert validation is still pending.", .8, 1.85, 11.7, 4.9, 20)

    s = slide("Before these can become thesis findings", "NEXT MILESTONE")
    add_text(s, "1  Verify OCR, section boundaries, dates and provenance against the PDFs.\n2  Review boilerplate, unusually long sections and repeated texts.\n3  Freeze a verified corpus; rerun both paths on identical inclusion rules.\n4  Use ordinance-grouped development / held-out evaluation before final selection.\n5  Expert panel: legal coherence, relevance, label appropriateness, temporal interpretation, overall quality (1–5 / N/A).\n6  Review any expert dimension below 3.00; regenerate the presentation from new evidence.", .8, 1.85, 11.7, 4.9, 20)

    for year in config["years"]:
        s = slide(f"{year} · both models, side by side", "YEAR-BY-YEAR EVIDENCE")
        picture(s, figures / f"year_{year}_comparison.png", y=1.7, h=4.7)
        add_text(s, "Topic numbers are local to each model: A/T0 is not B/T0. Bars show the top three topics, other assigned topics and outliers within this year's sample.", .8, 6.42, 11.7, .5, 14, GRAY)

    diagram = ROOT / "docs/reference_diagram.png"
    if diagram.exists():
        s = slide("Original supplied process diagram", "REFERENCE APPENDIX")
        picture(s, diagram, x=.8, y=1.65, w=5, h=5.1)
        add_text(s, "Source: Downloads/diagram.png\n\nImplementation clarification:\nc-TF-IDF needs class assignments. The lexical path therefore uses TF-IDF/SVD before clustering, then c-TF-IDF for topic representation.\n\nThe current deliverable is an exploratory modeling and reporting pipeline—not a completed expert-validated system.", 6.7, 1.9, 5.8, 4.9, 19)

    prs.save(out / "Ordinance_Two_Model_Initial_Results.pptx")
    summary = ["# Initial two-model results", "", config["interpretation"], "", "```csv", metrics.to_csv(index=False, lineterminator="\n"), "```",
               "", "## Artifacts", "- Editable PowerPoint: Ordinance_Two_Model_Initial_Results.pptx",
               "- figures/: parameter pictures, actual BERT trace, heatmaps and 10 year-by-year comparison PNGs",
               "- evidence/: source metrics, configuration, hashes and representative excerpts",
               "", "No claim of a winning model or thesis validation is made."]
    (out / "RESULTS.md").write_text("\n".join(summary), encoding="utf-8", newline="\n")
    print(f"Built {len(prs.slides)} slides and {len(list(figures.glob('*.png')))} PNG figures: {out}")


if __name__ == "__main__":
    main()
