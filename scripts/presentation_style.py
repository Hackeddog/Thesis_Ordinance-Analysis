"""Presentation layout helpers; no data or model assumptions."""
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import pandas as pd
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt
NAVY, TEAL, ORANGE, INK, GRAY = "10243A", "087F8C", "D47732", "18334A", "607487"


def pretty(value, digits=3):
    return "Not measured" if pd.isna(value) else f"{value:.{digits}f}"


def save_table(df, path, title, note=""):
    fig, ax = plt.subplots(figsize=(14, max(3.5, len(df) * .52 + 1.7)))
    ax.axis("off")
    ax.set_title(title, loc="left", fontsize=19, fontweight="bold", pad=20)
    table = ax.table(cellText=df.astype(str).values, colLabels=df.columns, loc="center", cellLoc="center")
    table.auto_set_font_size(False)
    table.set_fontsize(10)
    table.scale(1, 2)
    for (row, col), cell in table.get_celld().items():
        cell.set_edgecolor("white")
        cell.set_facecolor("#" + NAVY if row == 0 else ("#EEF4F7" if row % 2 else "#F8FAFC"))
        if row == 0:
            cell.set_text_props(color="white", weight="bold")
    fig.text(.02, .015, note, fontsize=9, color="#" + GRAY)
    fig.savefig(path, dpi=170, bbox_inches="tight")
    plt.close(fig)


def add_text(slide, text, x, y, w, h, size=18, color=INK, bold=False):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = line
        p.font.name = "Aptos"
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = RGBColor.from_string(color)
        p.space_after = Pt(12)
    return box


def picture(slide, path, x=0.7, y=1.65, w=11.9, h=4.85):
    from PIL import Image
    with Image.open(path) as image:
        ratio = image.width / image.height
    width, height = min(w, h * ratio), min(h, w / ratio)
    slide.shapes.add_picture(str(path), Inches(x + (w - width) / 2), Inches(y + (h - height) / 2),
                             width=Inches(width), height=Inches(height))
