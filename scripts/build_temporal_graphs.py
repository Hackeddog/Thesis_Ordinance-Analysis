"""Temporal-only figures with auditable draft titles derived from clustered keywords."""
from __future__ import annotations
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import textwrap

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
import numpy as np
import pandas as pd
from pptx import Presentation
from pptx.util import Inches

ROOT = Path(__file__).resolve().parents[1]
YEARS = list(range(2016, 2026))
NAMES = {"lexical": "Model A · Lexical", "semantic": "Model B · Legal-BERT"}
PALETTE = ["#087F8C", "#D47732", "#4768B1", "#A64D79", "#718C3B", "#7B61A8", "#9C6538", "#2F839A"]
RULES = [
    ("Disaster relief and calamity funding", "policy", "calamity calamities qrf quick fund assistance typhoon disaster"),
    ("Reconstruction and contingency appropriations", "policy", "reconstruction lump sum estimated calamities aside"),
    ("Temporary road closures and traffic control", "policy", "street streets closure temporary vehicular traffic road closed"),
    ("Event-related road closures", "policy", "rallies alley fiesta fairs celebrations temporarily closed"),
    ("Land-use zoning and reclassification", "policy", "zoning zone zones reclassification land agricultural amendment"),
    ("Amusement-tax exemptions", "policy", "amusement exemption proceeds exhibitions concerts operas oratorical"),
    ("Trade-event permit and fee exemptions", "policy", "expo trade exemption permit exhibitors payment"),
    ("Health workforce and midwifery support", "policy", "midwives midwife rhmpp placement healthdavao health"),
    ("Child protection and welfare", "policy", "child children abuse sexual protection"),
    ("Education and scholarship assistance", "policy", "scholarship scholarships ebsu semester students education tuition"),
    ("Single-use plastics regulation", "policy", "singleuse plastic plastics bags products"),
    ("Waste management and environmental protection", "policy", "waste garbage solid recycling disposal environment"),
    ("Public safety and CCTV systems", "policy", "cctv camera cameras surveillance security"),
    ("Cooperative financing and loan agreements", "policy", "cooperative multipurpose loan financing loans"),
    ("Public-health and medical services", "policy", "medical hospital health clinic healthcare patients"),
    ("Housing and community development", "policy", "housing homeowners settlement residential association"),
    ("Business licensing and local taxation", "policy", "business tax taxes license licenses licensing permits"),
    ("Agricultural and fishery support", "policy", "agriculture agricultural fishery fisheries farmers fishing"),
    ("Water supply and sanitation", "policy", "water sanitation sewerage drainage supply"),
    ("Traffic and parking regulation", "policy", "parking traffic vehicles vehicle transport road"),
    ("Gender and social-welfare programs", "policy", "women gender welfare disability disabilities senior"),
    ("Sports and community recreation", "policy", "sports games athletes recreation athletic"),
    ("Donation and property-transfer agreements", "governance", "donation donatton deed donor donee donated"),
    ("Mayoral contracting powers", "governance", "contract contracts powers chief executive enter behalf"),
    ("Agreements and mayoral authorization", "governance", "memorandum agreement moa authorizing entered"),
    ("Mayoral signing and legislative authority", "governance", "sign authority legislative behalf granted"),
    ("Government staffing and qualifications", "governance", "officer positions position training administrative experience relevant"),
    ("Public-fund utilization and audit controls", "governance", "utilization disbursements auditing coa dbm funds appropriated"),
    ("Council correspondence and legal endorsements", "procedural", "councilor atty respectfully secretary legal endorsement indorsement"),
    ("Council attendance and session records", "procedural", "councilor councibr leave jr sr counrilor counclbr"),
    ("Voting and enactment records", "procedural", "unanimous quorum vote eilacted eiacted members sanggunian"),
    ("Effectivity and publication clauses", "procedural", "effectmty effectivity immediately approval effect publication newspaper circulation"),
    ("Separability and validity clauses", "procedural", "separabiliw separability unconstitutional invalid continue force hereof"),
    ("Ordinance titles and amendment wording", "procedural", "known ordinance amending series title ordtnance sertes"),
]
OCR_MARKERS = {"ard", "thb", "ordinane", "gty", "fte", "fur", "sectioii", "sectioh", "aiid", "davm", "dayao", "sre", "ttre", "councibr", "councihr"}


def label_topic(model, topic, words, overrides):
    key = f"{model}:{topic}"
    words = [str(w).lower() for w in words]
    if key in overrides:
        chosen = overrides[key]
        if not set(chosen["required_terms"]).issubset(words):
            raise ValueError(f"Title override {key} does not match this run's clustered keywords")
        return chosen["title"], chosen["category"], "draft_keyword_excerpt_review", chosen["required_terms"]
    first = words[:8]
    if len(set(first) & OCR_MARKERS) >= 3 or sum(len(w) <= 2 for w in first) >= 5:
        return "OCR-heavy or unclear text cluster", "unclear", "keyword_quality_flag", [w for w in first if w in OCR_MARKERS]
    candidates = []
    for title, category, terms in RULES:
        hits = [w for w in words if w in set(terms.split())]
        if len(set(hits)) >= 2:
            score = sum(len(words) - words.index(w) for w in dict.fromkeys(hits))
            candidates.append((score, title, category, hits))
    if candidates:
        _, title, category, hits = max(candidates, key=lambda x: (x[0], x[1]))
        return title, category, "rank_weighted_keyword_rule", hits
    return "Unresolved keywords: " + ", ".join(words[:3]), "unclear", "keyword_fallback", words[:3]


def load_temporal(source, model, overrides):
    directory = Path(source) / model
    weighted = pd.read_csv(directory / "document_weighted_prevalence.csv")
    keywords = pd.read_csv(directory / "topic_keywords.csv")
    if weighted.duplicated(["year", "topic"]).any():
        raise ValueError("Duplicate topic/year cells")
    if set(weighted.year) != set(YEARS):
        raise ValueError("The complete study window must have observed documents")
    if weighted.groupby("year").documents_in_year.nunique().gt(1).any() or (weighted.documents_in_year <= 0).any():
        raise ValueError("Inconsistent yearly denominators")
    if (weighted.fractional_documents < 0).any() or not np.isfinite(weighted.share).all():
        raise ValueError("Invalid temporal weights")
    if not np.allclose(weighted.share, weighted.fractional_documents / weighted.documents_in_year):
        raise ValueError("Shares disagree with source numerators and denominators")
    if not np.allclose(weighted.groupby("year").share.sum(), 1):
        raise ValueError("Temporal shares including outliers must sum to one")
    totals = weighted.groupby("topic").fractional_documents.sum()
    matrix = weighted.pivot(index="topic", columns="year", values="share").reindex(columns=YEARS).fillna(0)
    catalog = []
    for topic in sorted(set(weighted.topic) - {-1}):
        words = keywords[keywords.topic.eq(topic)].sort_values("rank").term.tolist()
        if not words:
            raise ValueError(f"Missing keywords for topic {topic}")
        title, category, method, anchors = label_topic(model, topic, words, overrides)
        values = matrix.loc[topic]
        peak = values.max()
        peaks = [int(y) for y in YEARS if np.isclose(values[y], peak, rtol=0, atol=1e-12)]
        catalog.append({"model": model, "topic": int(topic), "title": title, "category": category,
                        "label_status": "draft_not_expert_validated", "label_method": method,
                        "matched_keywords": "; ".join(anchors), "clustered_words": "; ".join(words),
                        "fractional_documents_total": float(totals.loc[topic]),
                        "peak_years": "; ".join(map(str, peaks)), "peak_share_percent": peak * 100,
                        "share_2016_percent": values[2016] * 100, "share_2025_percent": values[2025] * 100,
                        "change_2016_to_2025_percentage_points": (values[2025] - values[2016]) * 100})
    return pd.DataFrame(catalog).sort_values(["fractional_documents_total", "topic"], ascending=[False, True]), matrix, weighted


def select_interpretable(catalog, n):
    # Do not merge similarly titled clusters; select the largest example of each title.
    return catalog[catalog.category.isin(["policy", "governance"])].drop_duplicates("title").head(n)


def topic_mix(matrix, selected):
    chosen = matrix.reindex(index=list(selected)).fillna(0)
    outliers = matrix.loc[-1] if -1 in matrix.index else pd.Series(0., index=YEARS)
    other = matrix.drop(index=list(selected) + [-1], errors="ignore").sum(axis=0)
    result = pd.concat([chosen, pd.DataFrame([other, outliers], index=["Other assigned topics", "Outliers"])]).T
    if not np.allclose(result.sum(axis=1), 1):
        raise ValueError("Grouped temporal shares do not conserve total document weight")
    return result


def build(source, output, titles, top_n=8):
    source, output, titles = Path(source), Path(output), Path(titles)
    if output.exists():
        raise ValueError("Output exists; choose a new directory to preserve the earlier figures")
    if not 1 <= top_n <= 8:
        raise ValueError("--top-n must be 1..8 for legible figures")
    config = json.loads(titles.read_text())
    loaded = {model: load_temporal(source, model, config["overrides"]) for model in NAMES}
    denominators = [value[2].groupby("year").documents_in_year.first().to_numpy() for value in loaded.values()]
    if not np.array_equal(denominators[0], denominators[1]):
        raise ValueError("Both models must describe the same yearly corpus")
    output.mkdir(parents=True)
    figures = output / "figures"; figures.mkdir()
    tables = output / "tables"; tables.mkdir()
    catalogs, summary_rows, figure_paths = [], [], []
    source_hashes = {}
    for model in NAMES:
        for filename in ["document_weighted_prevalence.csv", "topic_keywords.csv"]:
            path = source / model / filename
            source_hashes[f"{model}/{filename}"] = hashlib.sha256(path.read_bytes()).hexdigest()
    shared_limit = max(float(matrix.reindex(index=cat.head(top_n).topic).to_numpy().max()) * 100
                       for cat, matrix, _ in loaded.values())
    shared_limit = max(5, np.ceil(shared_limit / 5) * 5)
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 11, "axes.spines.top": False,
                         "axes.spines.right": False, "axes.titleweight": "bold"})
    pdf_path = output / "Temporal_Topic_Summaries.pdf"
    with PdfPages(pdf_path) as pdf:
        def save(fig, name):
            path = figures / name
            fig.savefig(path, dpi=190, facecolor="white")
            pdf.savefig(fig, facecolor="white")
            figure_paths.append(path)
            plt.close(fig)
        for model, (catalog, matrix, weighted) in loaded.items():
            catalogs.append(catalog)
            top = catalog.head(top_n)
            selected = select_interpretable(catalog, top_n)
            doc_counts = weighted.groupby("year").documents_in_year.first()
            catalog.to_csv(tables / f"{model}_topic_titles_and_temporal_summary.csv", index=False)
            long = weighted.merge(catalog[["topic", "title", "category", "label_status"]], on="topic", how="left")
            long.loc[long.topic.eq(-1), ["title", "category", "label_status"]] = ["Outliers", "unassigned", "not_a_substantive_topic"]
            long.to_csv(tables / f"{model}_labeled_temporal_data.csv", index=False)
            for selection_name, selection in [("leading", top), ("interpretable", selected)]:
                selection.assign(selection=selection_name).to_csv(tables / f"{model}_{selection_name}_topic_selection.csv", index=False)
                for row in selection.itertuples():
                    summary_rows.append({"model": model, "selection": selection_name, "topic": row.topic,
                        "title": row.title, "summary": f"Highest observed share: {row.peak_share_percent:.2f}% in {row.peak_years}; 2025 share: {row.share_2025_percent:.2f}%. Change from 2016: {row.change_2016_to_2025_percentage_points:+.2f} percentage points. Descriptive archive pattern, not a significance test."})

            # 1. Top-topic temporal heatmap, not a static topic-size chart.
            fig, ax = plt.subplots(figsize=(16, 9))
            values = matrix.loc[top.topic].to_numpy() * 100
            image = ax.imshow(values, aspect="auto", cmap="YlGnBu" if model == "lexical" else "YlOrBr", vmin=0, vmax=shared_limit)
            labels = [f"T{row.topic} · {textwrap.fill(row.title, 37)}" for row in top.itertuples()]
            ax.set_yticks(range(len(top)), labels, fontsize=11)
            ax.set_xticks(range(10), YEARS)
            ax.set_xlabel("Resolved study year")
            for i in range(len(top)):
                for j in range(10):
                    value = values[i,j]
                    ax.text(j, i, f"{value:.1f}", ha="center", va="center", fontsize=10,
                            color="white" if value > shared_limit*.6 else "#18334A")
            fig.colorbar(image, ax=ax, shrink=.82, label="Document-weighted share of year's corpus (%)")
            fig.suptitle(f"{NAMES[model]} | How leading topics vary over time", x=.035, ha="left", fontsize=22, weight="bold", color="#18334A")
            fig.text(.035, .905, "Top topics ranked by total fractional-document weight across 2016–2025. Draft titles summarize clustered keywords.", fontsize=11)
            fig.text(.035, .03, "Includes procedural/OCR-heavy clusters when they rank highly; do not interpret those as substantive policy themes.\nSame color scale across models. Topic IDs and titles do not imply cross-model equivalence.", fontsize=10, color="#607487")
            fig.subplots_adjust(left=.31, right=.93, top=.85, bottom=.15)
            save(fig, f"{model}_leading_topics_heatmap.png")

            # 2–3. Small multiples retain separate topic identities and their keyword evidence.
            for selection_name, selection in [("leading", top), ("interpretable", selected)]:
                nrows = max(1, (len(selection)+1)//2)
                fig, axes = plt.subplots(nrows, 2, figsize=(16, 3.05*nrows + 1.6), squeeze=False)
                max_share = max(2, np.ceil(matrix.reindex(index=selection.topic).to_numpy().max()*100/2)*2) if len(selection) else 2
                topic_colors = {t: PALETTE[i % len(PALETTE)] for i, t in enumerate(dict.fromkeys([*top.topic, *selected.topic]))}
                for ax, row in zip(axes.flat, selection.itertuples()):
                    shares = matrix.loc[row.topic] * 100
                    color = topic_colors[row.topic]
                    ax.plot(YEARS, shares.values, marker="o", color=color, linewidth=2, markersize=4)
                    ax.fill_between(YEARS, shares.values, alpha=.09, color=color)
                    ax.set_title(textwrap.fill(f"T{row.topic} · {row.title}", 48), loc="left", fontsize=12, pad=27)
                    evidence_words = row.matched_keywords or row.clustered_words
                    ax.text(0, 1.04, "Label evidence: " + ", ".join(evidence_words.split("; ")[:5]), transform=ax.transAxes, fontsize=9, color="#607487")
                    ax.set(ylim=(0, max_share*1.15), ylabel="Yearly share (%)")
                    ax.set_xticks(YEARS)
                    ax.tick_params(axis="x", labelsize=8)
                    ax.grid(axis="y", alpha=.18)
                    peak_year = int(shares.idxmax())
                    ax.annotate(f"{shares.max():.1f}%", (peak_year, shares.max()), xytext=(0, 8), textcoords="offset points", ha="center", fontsize=9, color=color)
                for ax in list(axes.flat)[len(selection):]:
                    ax.axis("off")
                descriptor = "Leading topic trajectories" if selection_name == "leading" else "Selected interpretable topic trajectories"
                fig.suptitle(f"{NAMES[model]} | {descriptor}", x=.045, ha="left", fontsize=22, weight="bold", color="#18334A")
                selection_note = "Top clusters by total document weight; procedural and unclear clusters remain visible." if selection_name == "leading" else "Policy/governance keyword labels only; largest cluster per repeated title. This is a disclosed subset, not merged topics."
                fig.text(.045, .938, selection_note + " Titles are drafts, not expert ratings.", fontsize=10)
                fig.text(.045, .016, "Document-weighted share = average within-document topic share in each year's included corpus; outliers remain in the denominator.\nLines connect annual observations only. Changes in this archive do not establish citywide policy growth, causation or statistical significance.", fontsize=10, color="#607487")
                fig.subplots_adjust(left=.07, right=.98, top=.865, bottom=.09, hspace=.92, wspace=.24)
                save(fig, f"{model}_{selection_name}_topic_trajectories.png")

            # 4. All weight is retained: displayed topics + other assigned topics + outliers.
            mix = topic_mix(matrix, top.topic)
            (mix * 100).to_csv(tables / f"{model}_yearly_mix_percent.csv", index_label="year", float_format="%.8f")
            fig, ax = plt.subplots(figsize=(16, 9))
            bottom = np.zeros(len(YEARS))
            titles_by_id = catalog.set_index("topic").title.to_dict()
            for i, column in enumerate(mix):
                color = PALETTE[i] if i < len(top) else ("#CBD5DF" if column == "Other assigned topics" else "#566573")
                label = f"T{column} · {titles_by_id[column]}" if isinstance(column, (int, np.integer)) else column
                ax.bar(YEARS, mix[column].values*100, bottom=bottom, color=color, width=.72,
                       label=textwrap.fill(label, 37), edgecolor="white", linewidth=.35)
                bottom += mix[column].values*100
            ax.set(ylim=(0, 100), ylabel="Document-weighted share of year's corpus (%)", xlabel="Resolved study year")
            ax.set_xticks(YEARS, [f"{year}\nN={int(doc_counts.loc[year])}" for year in YEARS], fontsize=10)
            ax.legend(loc="upper left", bbox_to_anchor=(1.01, 1), frameon=False, fontsize=10, labelspacing=1)
            fig.suptitle(f"{NAMES[model]} | Topic composition by year", x=.04, ha="left", fontsize=22, weight="bold", color="#18334A")
            fig.text(.04, .905, "Leading topics plus all remaining assigned topics and outliers; every year's column totals 100%.", fontsize=11)
            fig.text(.04, .035, "N = included source documents, not text units or all city ordinances. Each document contributes total weight 1.\nDraft descriptive titles; similar titles in different models do not establish an exact topic match.", fontsize=10, color="#607487")
            fig.subplots_adjust(left=.08, right=.64, top=.85, bottom=.17)
            save(fig, f"{model}_yearly_topic_composition.png")
    catalog = pd.concat(catalogs, ignore_index=True)
    catalog.to_csv(tables / "all_topic_titles.csv", index=False)
    pd.DataFrame(summary_rows).to_csv(tables / "selected_topic_temporal_summaries.csv", index=False)
    shutil_config = output / "title_rules_overrides.json"
    shutil_config.write_text(json.dumps(config, indent=2), encoding="utf-8")
    manifest = {"created_utc": datetime.now(timezone.utc).isoformat(), "source": str(source),
                "source_sha256": source_hashes, "title_config_sha256": hashlib.sha256(titles.read_bytes()).hexdigest(),
                "generator_sha256": hashlib.sha256(Path(__file__).read_bytes()).hexdigest(),
                "models_rerun": False, "years": YEARS, "top_n": top_n, "topic_titles": len(catalog),
                "figure_count": len(figure_paths), "label_status": config["status"],
                "ranking": "sum of fractional-document contributions across years, excluding topic -1",
                "denominator": "all included source documents in each year, including outlier contribution",
                "interpretable_selection": "policy/governance categories, highest-weight cluster per identical draft title; no merging",
                "yearly_documents": {str(y): int(n) for y, n in loaded['lexical'][2].groupby('year').documents_in_year.first().items()}}
    (output / "manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")
    # A graph-only slide deck: no model metrics, training screenshots or non-temporal charts.
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(16), Inches(10)
    prs.core_properties.title = "Temporal topic summaries — draft descriptive titles"
    from PIL import Image
    for path in figure_paths:
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        with Image.open(path) as image:
            ratio = image.width / image.height
        w, h = min(15.8, 9.8*ratio), min(9.8, 15.8/ratio)
        slide.shapes.add_picture(str(path), Inches((16-w)/2), Inches((10-h)/2), width=Inches(w), height=Inches(h))
    prs.save(output / "Temporal_Topic_Summaries.pptx")
    (output / "README.md").write_text(
        "# Temporal topic graphs\n\nEight temporal-only PNGs, an eight-page PDF, and a graph-only PowerPoint. "
        "These summarize existing initial-configuration results; no models were retrained.\n\n"
        "- Leading-topic heatmaps and trajectories retain procedural/OCR-heavy topics when they rank highly.\n"
        "- Interpretable trajectories are a disclosed policy/governance subset: the largest cluster per draft title, not a merging of clusters.\n"
        "- Yearly composition retains other assigned topics and outliers so every year totals 100%.\n"
        "- tables/all_topic_titles.csv lists all model-local topics and their keyword-based draft titles.\n"
        "- tables/selected_topic_temporal_summaries.csv records peak year(s), endpoint shares and descriptive change.\n\n"
        "Shares are document-weighted. Source documents, not all city ordinances, form the denominator. "
        "Titles are descriptive aids, not expert-approved categories. Equal titles across models do not establish alignment. "
        "Source verification does not make OCR/procedural clusters substantive policy themes.\n",
        encoding="utf-8", newline="\n")
    print(f"Created {len(figure_paths)} temporal figures, PDF/PPTX, and {len(catalog)} draft topic titles: {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source", type=Path, default=ROOT / "artifacts/clean_text_full/evidence")
    parser.add_argument("--output", type=Path, default=ROOT / "artifacts/temporal_topics")
    parser.add_argument("--titles", type=Path, default=ROOT / "configs/temporal_topic_titles.json")
    parser.add_argument("--top-n", type=int, default=8)
    args = parser.parse_args()
    build(args.source, args.output, args.titles, args.top_n)
